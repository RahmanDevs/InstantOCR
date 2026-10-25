import io
import json
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

class DocumentConverter:
    @staticmethod
    def to_txt(text: str) -> bytes:
        """Converts extracted text to plain text bytes."""
        return (text or "").encode("utf-8")

    @staticmethod
    def to_json(text: str, filename: str = "") -> bytes:
        """Converts extracted text to a structured JSON byte representation."""
        data = {
            "status": "success",
            "filename": filename,
            "extracted_text": text or "",
            "line_count": len((text or "").splitlines()),
            "character_count": len(text or ""),
            "word_count": len((text or "").split())
        }
        return json.dumps(data, indent=2, ensure_ascii=False).encode("utf-8")

    @staticmethod
    def to_csv(text: str) -> bytes:
        """Converts text lines into a structured CSV file."""
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(["Line", "Extracted_Text"])

        lines = (text or "").splitlines()
        for idx, line in enumerate(lines, start=1):
            if line.strip():
                writer.writerow([idx, line.strip()])

        return output.getvalue().encode("utf-8")

    @staticmethod
    def to_docx(text: str, title: str = "InstantOCR Extraction Result") -> bytes:
        """Creates a styled Microsoft Word (.docx) document."""
        doc = Document()

        # Set 1-inch standard margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Document Heading
        heading = doc.add_heading(title, level=1)
        if heading.runs:
            heading.runs[0].font.color.rgb = RGBColor(37, 99, 235)  # Primary blue

        doc.add_paragraph()  # Spacing

        lines = (text or "").splitlines()
        if not any(l.strip() for l in lines):
            doc.add_paragraph("(No text extracted from document)")
        else:
            for line in lines:
                if line.strip():
                    p = doc.add_paragraph(line.strip())
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.25

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def to_pdf(text: str, title: str = "InstantOCR Extraction Result") -> bytes:
        """Creates a styled Adobe PDF document using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45
        )
        styles = getSampleStyleSheet()

        body_style = ParagraphStyle(
            "OCRBody",
            parent=styles["Normal"],
            fontSize=10,
            leading=15,
            spaceAfter=6,
            textColor=colors.HexColor("#1e293b"),
            wordWrap="CJK"
        )

        title_style = ParagraphStyle(
            "OCRTitle",
            parent=styles["Title"],
            fontSize=16,
            leading=22,
            spaceAfter=14,
            textColor=colors.HexColor("#2563eb")
        )

        story = [
            Paragraph(f"<b>{title}</b>", title_style),
            Spacer(1, 10)
        ]

        def escape_xml(s: str) -> str:
            return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        lines = (text or "").splitlines()
        has_content = False
        for line in lines:
            if line.strip():
                has_content = True
                clean_line = escape_xml(line.strip())
                story.append(Paragraph(clean_line, body_style))

        if not has_content:
            story.append(Paragraph("<i>(No text detected in document)</i>", body_style))

        doc.build(story)
        return buffer.getvalue()

    @staticmethod
    def to_xlsx(text: str) -> bytes:
        """Creates an Excel (.xlsx) spreadsheet with auto-fitted columns."""
        lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
        if not lines:
            lines = ["(No text detected)"]

        df = pd.DataFrame({
            "Line": list(range(1, len(lines) + 1)),
            "Extracted Text": lines
        })

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="OCR Output")
            worksheet = writer.sheets["OCR Output"]

            # Auto-fit column width
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 100)

        return buffer.getvalue()

    @classmethod
    def convert(cls, text: str, format_type: str, filename: str = "") -> tuple[bytes, str, str]:
        """
        Converts text into the requested format.
        Returns: (file_bytes, content_type, file_extension)
        """
        fmt = (format_type or "txt").lower().strip()

        if fmt == "docx":
            return cls.to_docx(text, f"InstantOCR - {filename}" if filename else "InstantOCR Result"), \
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"

        elif fmt == "pdf":
            return cls.to_pdf(text, f"InstantOCR - {filename}" if filename else "InstantOCR Result"), \
                "application/pdf", "pdf"

        elif fmt == "xlsx":
            return cls.to_xlsx(text), \
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"

        elif fmt == "csv":
            return cls.to_csv(text), "text/csv; charset=utf-8", "csv"

        elif fmt == "json":
            return cls.to_json(text, filename), "application/json; charset=utf-8", "json"

        else:
            return cls.to_txt(text), "text/plain; charset=utf-8", "txt"
